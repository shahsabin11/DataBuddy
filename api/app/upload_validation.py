"""Validate upload bytes by content (not extension) and scan for unsafe cells."""

from __future__ import annotations

import csv
import io
import zipfile
from dataclasses import dataclass

from openpyxl import load_workbook

# Formula / CSV-injection prefixes (OWASP spreadsheet injection).
_INJECTION_PREFIXES = ("=", "+", "-", "@")

_OLE_MAGIC = b"\xd0\xcf\x11\xe0"
_ZIP_MAGIC = b"PK\x03\x04"


class UploadValidationError(Exception):
    def __init__(self, code: str, message: str):
        self.code = code
        self.message = message
        super().__init__(message)


@dataclass(frozen=True, slots=True)
class ValidatedUpload:
    kind: str  # csv | xlsx
    content_type: str
    extension: str
    row_count: int | None
    column_count: int | None


def assert_size_allowed(size: int, max_bytes: int) -> None:
    if size <= 0:
        raise UploadValidationError("empty_file", "The file is empty.")
    if size > max_bytes:
        mb = max_bytes // (1024 * 1024)
        raise UploadValidationError(
            "file_too_large",
            f"File exceeds the {mb} MB upload limit.",
        )


def _looks_like_text(sample: bytes) -> bool:
    if not sample:
        return False
    if b"\x00" in sample:
        return False
    # Reject high binary ratio in the sniff window.
    weird = sum(1 for b in sample if b < 9 or (13 < b < 32 and b not in (9, 10, 13)))
    return weird / max(len(sample), 1) < 0.3


def _cell_is_injection(value: object) -> bool:
    if value is None:
        return False
    if not isinstance(value, str):
        return False
    stripped = value.lstrip("\t\r\n ")
    if not stripped:
        return False
    return stripped[0] in _INJECTION_PREFIXES


def _scan_csv(data: bytes) -> tuple[int, int]:
    try:
        text = data.decode("utf-8-sig")
    except UnicodeDecodeError as exc:
        raise UploadValidationError(
            "invalid_encoding",
            "CSV must be UTF-8 encoded.",
        ) from exc

    reader = csv.reader(io.StringIO(text))
    rows = 0
    cols = 0
    for row in reader:
        rows += 1
        cols = max(cols, len(row))
        for cell in row:
            if _cell_is_injection(cell):
                raise UploadValidationError(
                    "formula_injection",
                    "File contains a cell that looks like a spreadsheet formula "
                    "(starts with =, +, -, or @). Remove it and upload again.",
                )
    if rows == 0:
        raise UploadValidationError("empty_file", "The CSV has no rows.")
    return rows, cols


def _scan_xlsx(data: bytes) -> tuple[int, int]:
    if not data.startswith(_ZIP_MAGIC):
        raise UploadValidationError(
            "invalid_type",
            "File is not a valid XLSX workbook.",
        )
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as zf:
            names = set(zf.namelist())
            if "[Content_Types].xml" not in names or not any(
                n.startswith("xl/") for n in names
            ):
                raise UploadValidationError(
                    "invalid_type",
                    "File is not a valid XLSX workbook.",
                )
            # Macro-enabled or embedded VBA / macro sheets.
            lower_names = {n.lower() for n in names}
            if any(
                n.endswith("vbaproject.bin")
                or "/macrosheets/" in n
                or n.endswith(".xlsm")
                for n in lower_names
            ):
                raise UploadValidationError(
                    "macros_forbidden",
                    "Spreadsheets with macros are not allowed.",
                )
    except zipfile.BadZipFile as exc:
        raise UploadValidationError(
            "invalid_type",
            "File is not a valid XLSX workbook.",
        ) from exc

    try:
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=False)
    except Exception as exc:  # openpyxl raises variously for corrupt files
        raise UploadValidationError(
            "invalid_type",
            "Could not read the XLSX workbook.",
        ) from exc

    try:
        rows = 0
        cols = 0
        for sheet in wb.worksheets:
            for row in sheet.iter_rows(values_only=False):
                rows += 1
                cols = max(cols, len(row))
                for cell in row:
                    # Reject stored formulas and injection-looking values.
                    if getattr(cell, "data_type", None) == "f" or (
                        isinstance(cell.value, str) and cell.value.startswith("=")
                    ):
                        raise UploadValidationError(
                            "formula_injection",
                            "File contains spreadsheet formulas. "
                            "Upload values only (CSV or formula-free XLSX).",
                        )
                    if _cell_is_injection(cell.value):
                        raise UploadValidationError(
                            "formula_injection",
                            "File contains a cell that looks like a spreadsheet "
                            "formula (starts with =, +, -, or @).",
                        )
        if rows == 0:
            raise UploadValidationError("empty_file", "The workbook has no rows.")
        return rows, cols
    finally:
        wb.close()


def validate_upload(data: bytes, *, filename: str, max_bytes: int) -> ValidatedUpload:
    """
    Detect CSV vs XLSX from content. Extension is only a hint for naming.
    """
    assert_size_allowed(len(data), max_bytes)

    if data.startswith(_OLE_MAGIC):
        raise UploadValidationError(
            "invalid_type",
            "Legacy .xls / OLE files are not supported. Export as CSV or XLSX.",
        )

    name = (filename or "").lower()
    sample = data[:8192]

    if data.startswith(_ZIP_MAGIC):
        rows, cols = _scan_xlsx(data)
        return ValidatedUpload(
            kind="xlsx",
            content_type=(
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
            ),
            extension=".xlsx",
            row_count=rows,
            column_count=cols,
        )

    if _looks_like_text(sample) or name.endswith(".csv"):
        if not _looks_like_text(sample):
            raise UploadValidationError(
                "invalid_type",
                "Only CSV and XLSX uploads are supported.",
            )
        rows, cols = _scan_csv(data)
        return ValidatedUpload(
            kind="csv",
            content_type="text/csv",
            extension=".csv",
            row_count=rows,
            column_count=cols,
        )

    raise UploadValidationError(
        "invalid_type",
        "Only CSV and XLSX uploads are supported.",
    )
